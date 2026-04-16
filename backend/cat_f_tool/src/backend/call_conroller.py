import logging
import threading
from datetime import datetime
import os.path
import time
from queue import Queue
from typing import List

from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.workbook import Workbook

from backend.cat_f_tool.src.data.device_info      import DeviceInfo
from backend.cat_f_tool.src.data.dir_paths         import Paths
from backend.cat_f_tool.src.data.test_credentials  import TestCred
from backend.cat_f_tool.src.data.constants         import VERSION, COLUMN_NAMES

from backend.cat_f_tool.src.backend.helper import (
    dial_number, receive_call, end_call, get_call_state,
    check_network_registration_status, check_device_disconnection,
    wait_for_ring_on_mt, check_mo_dialing, answer_call_on_mt,
    setup_mt_for_incoming_call,
    collect_nw_params,                 # NEW — NW parameter collection
)
from backend.cat_f_tool.src.frontend.connection_break import reconnect
from backend.cat_f_tool.src.frontend.progress_check   import TestProgress

from backend.cat_f_tool.src.utility.logger import CustomLogger
from backend.cat_f_tool.src.utility.utils  import create_folders, send_at_command


class CallTester:
    def __init__(self, home, progress_frame,
                 resume_event: threading.Event,
                 pause_event: threading.Event,
                 end_test_event: threading.Event):
        self.resume_event   = resume_event
        self.pause_event    = pause_event
        self.end_event      = end_test_event
        self.call_duration  = None
        self.comment        = None
        self.result         = None
        self.home           = home
        self.progress_frame = progress_frame
        self._wb_lock       = threading.Lock()   # guards concurrent workbook saves

        self.logger = CustomLogger("backend", Paths.latest_log_dir_path)
        self.logger.log("Backend code initiated to Start the Call Testing.")

        self.test_type  = TestCred.test_type
        self.iterations = TestCred.iterations
        self.duration   = TestCred.call_duration
        self.delay      = TestCred.call_delay

        # MO calls  = MT device (dev2) calls MO device (dev1)
        # MT calls  = MO device (dev1) calls MT device (dev2)
        # MO-MT     = phase1 MO->MT then phase2 MT->MO
        if self.test_type == "MO calls":
            self.mo_sw   = DeviceInfo.dev2["Software Version"]
            self.mo_imei = DeviceInfo.dev2["IMEI"]
            self.mo_port = DeviceInfo.dev2["AT Port"]
            self.mo_no   = DeviceInfo.dev2["Phone Number"]
            self.mt_sw   = DeviceInfo.dev1["Software Version"]
            self.mt_imei = DeviceInfo.dev1["IMEI"]
            self.mt_port = DeviceInfo.dev1["AT Port"]
            self.mt_no   = DeviceInfo.dev1["Phone Number"]
        if self.test_type == "MT calls":
            self.mo_sw   = DeviceInfo.dev1["Software Version"]
            self.mo_imei = DeviceInfo.dev1["IMEI"]
            self.mo_port = DeviceInfo.dev1["AT Port"]
            self.mo_no   = DeviceInfo.dev1["Phone Number"]
            self.mt_sw   = DeviceInfo.dev2["Software Version"]
            self.mt_imei = DeviceInfo.dev2["IMEI"]
            self.mt_port = DeviceInfo.dev2["AT Port"]
            self.mt_no   = DeviceInfo.dev2["Phone Number"]
        if self.test_type == "MO-MT calls":
            self.mo_sw   = DeviceInfo.dev1["Software Version"]
            self.mo_imei = DeviceInfo.dev1["IMEI"]
            self.mo_port = DeviceInfo.dev1["AT Port"]
            self.mo_no   = DeviceInfo.dev1["Phone Number"]
            self.mt_sw   = DeviceInfo.dev2["Software Version"]
            self.mt_imei = DeviceInfo.dev2["IMEI"]
            self.mt_port = DeviceInfo.dev2["AT Port"]
            self.mt_no   = DeviceInfo.dev2["Phone Number"]

        # Report workbook
        self.bold_font = Font(bold=True)
        self.report_file_path = os.path.join(
            Paths.latest_report_dir_path,
            f'CAT_report_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        )
        self.workbook      = Workbook()
        self.current_sheet = self.workbook.active
        self.current_sheet.title = self.test_type if self.test_type != "MO-MT calls" else "MO calls"
        self._write_1st_line()
        self._write_device_info()
        self._write_report_columns_names(self.current_sheet, COLUMN_NAMES, 8)

    def _check_pause(self):
        if self.pause_event.isSet():
            self.pause_event.clear()
            self.logger.log(message="Automation Paused.", level=logging.INFO)
            # End any ongoing call on both sides so devices are idle while paused
            end_call(self.mo_port, self.logger)
            end_call(self.mt_port, self.logger)
            self.resume_event.wait()
            self.resume_event.clear()
            self.logger.log(message="Automation Resumed.", level=logging.INFO)

    def _write_1st_line(self):
        first_line_text = f"CAT Tool (version- {VERSION}) report"
        self.current_sheet.merge_cells(start_row=1, start_column=2, end_row=1, end_column=5)
        self.current_sheet.cell(row=1, column=2).value = first_line_text
        self.current_sheet.cell(row=1, column=2).font  = self.bold_font
        self._save_workbook()

    def _write_device_info(self):
        self.current_sheet.cell(row=3, column=2).value = "MO"
        self.current_sheet.cell(row=3, column=2).font  = self.bold_font
        self.current_sheet.cell(row=3, column=3).value = "MT"
        self.current_sheet.cell(row=3, column=3).font  = self.bold_font
        self.current_sheet.cell(row=4, column=1).value = "Software Version"
        self.current_sheet.cell(row=4, column=1).font  = self.bold_font
        self.current_sheet.cell(row=4, column=2).value = self.mo_sw
        self.current_sheet.cell(row=4, column=3).value = self.mt_sw
        self.current_sheet.cell(row=5, column=1).value = "IMEI"
        self.current_sheet.cell(row=5, column=1).font  = self.bold_font
        self.current_sheet.cell(row=5, column=2).value = self.mo_imei
        self.current_sheet.cell(row=5, column=3).value = self.mt_imei
        self.current_sheet.cell(row=6, column=1).value = "Phone Number"
        self.current_sheet.cell(row=6, column=1).font  = self.bold_font
        self.current_sheet.cell(row=6, column=2).value = self.mo_no
        self.current_sheet.cell(row=6, column=3).value = self.mt_no
        self._save_workbook()

    def _switch_mo_mt(self):
        self.mo_sw   = DeviceInfo.dev2["Software Version"]
        self.mo_imei = DeviceInfo.dev2["IMEI"]
        self.mo_port = DeviceInfo.dev2["AT Port"]
        self.mo_no   = DeviceInfo.dev2["Phone Number"]
        self.mt_sw   = DeviceInfo.dev1["Software Version"]
        self.mt_imei = DeviceInfo.dev1["IMEI"]
        self.mt_port = DeviceInfo.dev1["AT Port"]
        self.mt_no   = DeviceInfo.dev1["Phone Number"]

    def _write_report_columns_names(self, worksheet, columns: List[str], row: int):
        from openpyxl.styles import Alignment as XlAlign
        from openpyxl.utils import get_column_letter

        bg_color   = PatternFill(start_color="B0E0E6", end_color="B0E0E6", fill_type="solid")
        mo_color   = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
        mt_color   = PatternFill(start_color="FCE5CD", end_color="FCE5CD", fill_type="solid")

        # Column widths keyed on header name
        _widths = {
            "Iteration": 10, "Date": 12, "Start Time": 12, "End Time": 12,
            "Duration": 10, "Status": 10, "Comment": 22,
            "MO_MCC": 9,  "MO_MNC": 9,  "MO_Operator": 18,
            "MO_LAC_TAC": 12, "MO_Cell_ID": 12,
            "MO_RSSI_dBm": 13, "MO_BER": 8,
            "MO_RSRP_dBm": 13, "MO_RSRQ_dB": 12,
            "MO_RAT": 10,
            "MT_MCC": 9,  "MT_MNC": 9,  "MT_Operator": 18,
            "MT_LAC_TAC": 12, "MT_Cell_ID": 12,
            "MT_RSSI_dBm": 13, "MT_BER": 8,
            "MT_RSRP_dBm": 13, "MT_RSRQ_dB": 12,
            "MT_RAT": 10,
        }

        for col_idx, column_name in enumerate(columns, 1):
            cell       = worksheet.cell(row=row, column=col_idx, value=column_name)
            cell.font  = self.bold_font
            # MO columns → green tint, MT columns → orange tint, core → blue
            if column_name.startswith("MO_"):
                cell.fill = mo_color
            elif column_name.startswith("MT_"):
                cell.fill = mt_color
            else:
                cell.fill = bg_color
            cell.alignment = XlAlign(horizontal="center", vertical="center", wrap_text=True)
            col_letter = get_column_letter(col_idx)
            worksheet.column_dimensions[col_letter].width = _widths.get(column_name, 12)

        self._save_workbook()

    def _new_sheet_for_mt_calls(self):
        self._switch_mo_mt()
        self.current_sheet       = self.workbook.create_sheet()
        self.current_sheet.title = "MT calls"
        self._write_1st_line()
        self._write_device_info()
        self._write_report_columns_names(self.current_sheet, COLUMN_NAMES, 8)
        self._save_workbook()

    def _save_workbook(self):
        with self._wb_lock:
            self.workbook.save(self.report_file_path)

    def _update_report(self, itr, date, start_time, end_time, duration, result, comment,
                       nw_mo=None, nw_mt=None):
        """Append one row to the active sheet — core columns + NW params for MO and MT."""
        _empty = {
            "MCC": "N/A", "MNC": "N/A", "Operator": "N/A",
            "LAC_TAC": "N/A", "Cell_ID": "N/A",
            "RSSI_dBm": "N/A", "BER": "N/A",
            "RSRP_dBm": "N/A", "RSRQ_dB": "N/A",
            "RAT": "N/A",
        }
        mo = nw_mo or _empty
        mt = nw_mt or _empty
        row = [
            itr, date, start_time, end_time, duration, result, comment,
            mo["MCC"],      mo["MNC"],      mo["Operator"],
            mo["LAC_TAC"],  mo["Cell_ID"],
            mo["RSSI_dBm"], mo["BER"],
            mo["RSRP_dBm"], mo["RSRQ_dB"],
            mo["RAT"],
            mt["MCC"],      mt["MNC"],      mt["Operator"],
            mt["LAC_TAC"],  mt["Cell_ID"],
            mt["RSSI_dBm"], mt["BER"],
            mt["RSRP_dBm"], mt["RSRQ_dB"],
            mt["RAT"],
        ]
        self.current_sheet.append(row)
        self._save_workbook()

    def _make_devices_idle(self):
        """Send ATH to both ports to ensure they are in IDLE state."""
        self.logger.log("Putting devices into IDLE state...")
        end_call(self.mo_port, self.logger)
        end_call(self.mt_port, self.logger)
        time.sleep(1)

    def perform_call(self, current_itr):
        """
        Performs one iteration of MO→MT call testing.

        Designed to work correctly with both:
          - Feature phones that support AT+CLCC (standard)
          - 2G/Unisoc feature phones that only emit RING URCs (no AT+CLCC)

        Flow:
          1. Check both ports are connected
          2. Put devices into IDLE
          3. Dial from MO port
          4. Give MO 2s to start the call setup
          5. Check MO is dialing (AT+CLCC if supported; else trust ATD result)
          6. Listen on MT port directly for RING/+CSYS URC (up to 30s)
          7. Answer the call on MT with ATA
          8. Hold call for configured duration
          9. End call on both sides
         10. Write report row
        """
        date       = datetime.now().strftime("%d/%m/%Y")
        start_time = datetime.now().strftime("%H:%M:%S")
        end_time   = "00:00:00"
        self.call_duration = "00"

        self.logger.log(f"{'='*50}")
        self.logger.log(f"Iteration {current_itr} START")
        self.logger.log(f"  MO port: {self.mo_port}  |  MT port: {self.mt_port}")
        self.logger.log(f"  Dialing: {self.mt_no}  |  Duration: {self.duration}s")

        # ── Step 1: Disconnection check ───────────────────────────────────────
        status, _ = check_device_disconnection(self.mo_port, self.mt_port, self.logger)
        if not status:
            raise Exception("Device disconnected!")

        # ── Step 2: Put both devices into IDLE ────────────────────────────────
        self._make_devices_idle()

        # Step 2b: Setup MT BEFORE dialing - send AT+CRC=1/AT+CLIP=1
        # and keep the port open so settings survive until the call arrives
        mt_ser = setup_mt_for_incoming_call(self.mt_port, self.logger)

        # Step 3: Dial from MO
        self.logger.log(f"Dialing {self.mt_no} from MO port {self.mo_port}...")
        dial_success = dial_number(self.mo_port, self.mt_no, self.logger)
        if not dial_success:
            self.logger.log(f"Dial failed for {self.mt_no} on {self.mo_port}", level=logging.WARNING)
            if mt_ser: mt_ser.close()
            # Check if dial failure is due to OOS before marking plain Fail
            oos_result = self._check_and_handle_oos(current_itr, date, start_time, "Dial failed")
            if oos_result is not None:
                return oos_result
            self._update_report(current_itr, date, start_time, "NA", "00", "Fail", "MO dial failed")
            return

        # Step 4: Brief pause for network call setup
        self.logger.log("Dial accepted. Waiting 2s for call setup...")
        time.sleep(2)

        # Step 5: Confirm MO is dialing
        mo_dialing = check_mo_dialing(self.mo_port, self.logger)
        if not mo_dialing:
            self.logger.log("MO not dialing after ATD - marking Fail.")
            if mt_ser: mt_ser.close()
            end_call(self.mo_port, self.logger)
            # Check if failure is due to OOS
            oos_result = self._check_and_handle_oos(current_itr, date, start_time, "MO failed to establish call")
            if oos_result is not None:
                return oos_result
            self._update_report(current_itr, date, start_time, "NA", "00",
                                "Fail", "MO failed to establish call")
            return

        # Steps 6+7: listen on ALL HMD ports for RING then answer via ATA
        # wait_for_ring_on_mt monitors sibling COM ports in case RING URCs
        # arrive on a different port than the AT command port (COM9)
        self.logger.log("Listening for RING on all HMD ports and answering...")
        answered = wait_for_ring_on_mt(
            self.mt_port, timeout_sec=30, logger=self.logger, ser=mt_ser
        )
        if mt_ser:
            try: mt_ser.close()
            except Exception: pass
        if not answered:
            self.logger.log("MT did not ring or answer within 30s.")
            end_call(self.mo_port, self.logger)
            end_call(self.mt_port, self.logger)
            # Check if ring timeout is due to OOS
            oos_result = self._check_and_handle_oos(current_itr, date, start_time, "MT did not ring or answer")
            if oos_result is not None:
                return oos_result
            self._update_report(current_itr, date, start_time, "NA", "00",
                                "Fail", "MT did not ring or answer")
            return

        # ── Step 8: Record call start, hold for configured duration ───────────
        call_start = datetime.now()
        start_time = call_start.strftime("%H:%M:%S")
        hold_secs  = int(self.duration)
        self.logger.log(f"Call connected. Holding for {hold_secs}s...")

        # Interruptible hold — tracks wall-clock time so serial AT reads for
        # OOS checks never inflate the configured call duration.
        hold_deadline    = time.monotonic() + hold_secs
        paused_mid_call  = False
        _oos_check_every = 5
        _last_oos_check  = time.monotonic()

        while time.monotonic() < hold_deadline:
            time.sleep(1)

            # ── Pause check ───────────────────────────────────────────────────
            if self.pause_event.isSet():
                self.pause_event.clear()
                self.logger.log("Pause detected mid-call — ending call now.")
                end_call(self.mo_port, self.logger)
                end_call(self.mt_port, self.logger)
                self.logger.log(message="Automation Paused.", level=logging.INFO)
                self.resume_event.wait()
                self.resume_event.clear()
                self.logger.log(message="Automation Resumed.", level=logging.INFO)
                paused_mid_call = True
                break

            # ── OOS CHECK DURING ACTIVE CALL (every 5 s, not every tick) ─────
            now = time.monotonic()
            if now - _last_oos_check >= _oos_check_every:
                _last_oos_check = now
                mo_registered = check_network_registration_status(self.mo_port)
                mt_registered = check_network_registration_status(self.mt_port)
                if not mo_registered or not mt_registered:
                    still_connected, _ = check_device_disconnection(self.mo_port, self.mt_port, self.logger)
                    if not still_connected:
                        raise Exception("Device disconnected!")

                    oos_ports = []
                    if not mo_registered:
                        oos_ports.append(f"MO ({self.mo_port})")
                    if not mt_registered:
                        oos_ports.append(f"MT ({self.mt_port})")

                    self.logger.log(f"⚠️ OOS DURING CALL: {oos_ports} — ending call and waiting for network return")
                    end_call(self.mo_port, self.logger)
                    end_call(self.mt_port, self.logger)
                    end_time = datetime.now().strftime("%H:%M:%S")

                    oos_result = self._wait_for_network_return(oos_ports, current_itr, date, start_time, end_time)
                    return oos_result

        # ── Step 9: End call on both sides ────────────────────────────────────
        if not paused_mid_call:
            end_call(self.mo_port, self.logger)
            time.sleep(0.5)
            end_call(self.mt_port, self.logger)

        end_time = datetime.now().strftime("%H:%M:%S")

        if paused_mid_call:
            self.logger.log(f"Iteration {current_itr} interrupted by pause — will retry.")
            self._update_report(current_itr, date, start_time, end_time,
                                "00", "Paused", "Call stopped by pause — will retry")
            return "retry"

        self.call_duration = self.duration
        self.logger.log(f"Iteration {current_itr} PASS — call held for {hold_secs}s")

        # Collect NW params for both devices in PARALLEL — cuts collection time
        # from ~7s (sequential) to ~4s (concurrent) so the delay is closer to exact.
        self.logger.log("Collecting NW parameters from MO and MT devices...")
        nw_mo_result = [None]
        nw_mt_result = [None]

        def _collect_mo():
            nw_mo_result[0] = collect_nw_params(self.mo_port, self.logger)

        def _collect_mt():
            nw_mt_result[0] = collect_nw_params(self.mt_port, self.logger)

        t_mo = threading.Thread(target=_collect_mo, daemon=True)
        t_mt = threading.Thread(target=_collect_mt, daemon=True)
        t_mo.start()
        t_mt.start()
        t_mo.join()
        t_mt.join()

        nw_mo = nw_mo_result[0]
        nw_mt = nw_mt_result[0]
        self.logger.log(f"MO NW params: {nw_mo}")
        self.logger.log(f"MT NW params: {nw_mt}")
        self._update_report(current_itr, date, start_time, end_time,
                            self.call_duration, "Pass", "",
                            nw_mo=nw_mo, nw_mt=nw_mt)

    def _check_and_handle_oos(self, current_itr, date, start_time, original_reason):
        """
        Called after a call attempt fails. Checks if the failure was caused by OOS.
        If OOS is detected, logs it, notifies the UI, waits for network return,
        and returns "oos_retry" so the iteration is retried.
        Returns None if the failure was NOT due to OOS.
        """
        mo_registered = check_network_registration_status(self.mo_port)
        mt_registered = check_network_registration_status(self.mt_port)
        if mo_registered and mt_registered:
            return None   # Not OOS — caller handles normal failure

        # Confirm devices are still physically connected
        still_connected, _ = check_device_disconnection(self.mo_port, self.mt_port, self.logger)
        if not still_connected:
            return None   # Physical disconnect — let caller raise exception

        oos_ports = []
        if not mo_registered:
            oos_ports.append(f"MO ({self.mo_port})")
        if not mt_registered:
            oos_ports.append(f"MT ({self.mt_port})")

        self.logger.log(f"⚠️ OOS DETECTED (call failed due to no network): {oos_ports}")
        end_time = datetime.now().strftime("%H:%M:%S")
        return self._wait_for_network_return(oos_ports, current_itr, date, start_time, end_time)

    def _wait_for_network_return(self, oos_ports, current_itr, date, start_time, end_time):
        """
        Signals the UI with OOS alert, blocks until both devices regain network
        (or test is stopped), then returns the appropriate sentinel:
          "oos_retry"  — devices back online, run_test should retry this iteration
          "oos_abort"  — test was stopped during the wait
        Also writes an OOS row into the report.
        """
        from backend.cat_f_tool.service import cat_f_state

        self.logger.log(f"📵 OOS devices: {', '.join(oos_ports)} — pausing test, waiting for service...")

        # Write OOS row in report
        self._update_report(current_itr, date, start_time, end_time,
                            "00", "OOS", f"Device in OOS zone: {', '.join(oos_ports)}")

        # Signal the UI
        cat_f_state["oos_detected"] = True
        cat_f_state["oos_devices"]  = oos_ports
        my_test_id = cat_f_state.get("test_id")

        from backend.cat_f_tool.service import _stop_event

        # Poll until both devices are back in service
        while True:
            time.sleep(5)

            # Honour stop request
            if _stop_event.is_set():
                self.logger.log("⏹️ Test stopped during OOS wait")
                cat_f_state["oos_detected"] = False
                cat_f_state["oos_devices"]  = []
                return "oos_abort"

            # New test started?
            if cat_f_state.get("test_id") != my_test_id:
                self.logger.log("⚠️ New test started during OOS wait — aborting")
                return "oos_abort"

            # Check registration on both ports
            mo_ok = check_network_registration_status(self.mo_port)
            mt_ok = check_network_registration_status(self.mt_port)

            if mo_ok and mt_ok:
                self.logger.log("✅ Both devices back in service — retrying iteration")
                cat_f_state["oos_detected"] = False
                cat_f_state["oos_devices"]  = []
                return "oos_retry"
            else:
                still_oos = []
                if not mo_ok:
                    still_oos.append(f"MO ({self.mo_port})")
                if not mt_ok:
                    still_oos.append(f"MT ({self.mt_port})")
                self.logger.log(f"⏳ Still OOS: {', '.join(still_oos)} — waiting...")

    def _handle_disconnect(self, devices: list):
        """
        Called when a device disconnect is detected during run_test.
        - Ends any active call on available ports
        - Sets cat_f_state so the UI shows the disconnect popup and paused badge
        - Blocks until both ports reappear (or stop is requested)
        - Clears cat_f_state and returns so run_test can retry the same iteration
        Returns True  → device back, retry same iteration
        Returns False → stop was requested, abort run_test
        """
        from backend.cat_f_tool.service import cat_f_state as _svc, _stop_event

        # End any call on whichever port is still alive
        for port in [self.mo_port, self.mt_port]:
            if port not in devices:
                try:
                    end_call(port, self.logger)
                except Exception:
                    pass

        self.logger.log(
            f"⚠️ Device disconnected: {', '.join(devices)} — "
            f"test paused, waiting for reconnect...",
            level=logging.WARNING
        )

        # Signal UI
        _svc["device_disconnected"] = True
        _svc["disconnected_device"] = ", ".join(devices)

        # Wait until all ports are back — silent polling, no repeated log spam
        while True:
            if _stop_event.is_set():
                self.logger.log("⏹️ Test stopped during disconnect wait.")
                _svc["device_disconnected"] = False
                _svc["disconnected_device"] = None
                return False

            reconnect_status = reconnect(self.home, devices, 20)

            # Check silently — skip the logger so "Device(s) not found" doesn't
            # repeat on every poll tick while we are already showing the popup
            from serial.tools import list_ports as _lp
            available = [p.device for p in _lp.comports()]
            still_missing = [p for p in [self.mo_port, self.mt_port] if p not in available]

            if reconnect_status and not still_missing:
                self.logger.log("✅ Device reconnected — resuming from same iteration")
                _svc["device_disconnected"] = False
                _svc["disconnected_device"] = None
                self._make_devices_idle()
                return True
            # No log here — we already told the user once above

    def run_test(self, queue: Queue):
        self.logger.log("run_test started")
        progress    = TestProgress(self.progress_frame, self.iterations)
        current_itr = 1

        while current_itr <= self.iterations:
            try:
                self._check_pause()
                status, _ = check_device_disconnection(self.mo_port, self.mt_port, self.logger)
                if not status:
                    raise Exception("Device disconnected!")
                result = self.perform_call(current_itr)
                if result == "retry":
                    self.logger.log(f"Retrying iteration {current_itr} after resume.")
                    continue
                if result == "oos_retry":
                    self.logger.log(f"Retrying iteration {current_itr} after OOS recovery.")
                    continue
                if result == "oos_abort":
                    self.logger.log("OOS abort — ending run_test.")
                    return
                time.sleep(self.delay)
                progress.update_progress()
                queue.put(current_itr)
                current_itr += 1
                self._save_workbook()
                if self.end_event.is_set():
                    return
            except Exception as e:
                status, devices = check_device_disconnection(self.mo_port, self.mt_port, self.logger)
                if not status:
                    # Device physically disconnected — pause, wait, retry same iteration
                    resumed = self._handle_disconnect(devices)
                    if not resumed:
                        return   # stop was requested during wait
                    # current_itr NOT incremented — retry same iteration
                else:
                    self.logger.log(f"System error in iteration {current_itr}: {e}", level=logging.ERROR)
                    self._update_report(
                        current_itr,
                        datetime.now().strftime("%d/%m/%Y"),
                        datetime.now().strftime("%H:%M:%S"),
                        "NA", "00", "Fail", f"System error: {str(e)[:80]}"
                    )
                    current_itr += 1

        # ── MO-MT calls: second pass with roles swapped ────────────────────
        if self.test_type == "MO-MT calls":
            self._new_sheet_for_mt_calls()
            self.logger.log("Starting MT phase of MO-MT calls test")
            progress    = TestProgress(self.progress_frame, self.iterations)
            current_itr = 1
            while current_itr <= self.iterations:
                try:
                    self._check_pause()
                    status, _ = check_device_disconnection(self.mo_port, self.mt_port, self.logger)
                    if not status:
                        raise Exception("Device disconnected!")
                    result = self.perform_call(current_itr)
                    if result == "retry":
                        self.logger.log(f"Retrying MT iteration {current_itr} after resume.")
                        continue
                    if result == "oos_retry":
                        self.logger.log(f"Retrying MT iteration {current_itr} after OOS recovery.")
                        continue
                    if result == "oos_abort":
                        self.logger.log("OOS abort during MT phase — ending run_test.")
                        return
                    time.sleep(self.delay)
                    progress.update_progress()
                    queue.put(current_itr)
                    current_itr += 1
                    self._save_workbook()
                    if self.end_event.is_set():
                        return
                except Exception as e:
                    status, devices = check_device_disconnection(self.mo_port, self.mt_port, self.logger)
                    if not status:
                        # Device physically disconnected — pause, wait, retry same iteration
                        resumed = self._handle_disconnect(devices)
                        if not resumed:
                            return   # stop was requested during wait
                        # current_itr NOT incremented — retry same MT iteration
                    else:
                        self.logger.log(f"System error in MT iteration {current_itr}: {e}", level=logging.ERROR)
                        self._update_report(
                            current_itr,
                            datetime.now().strftime("%d/%m/%Y"),
                            datetime.now().strftime("%H:%M:%S"),
                            "NA", "00", "Fail", f"System error: {str(e)[:80]}"
                        )
                        current_itr += 1


def run_call_testing(test: CallTester, queue):
    test.run_test(queue)